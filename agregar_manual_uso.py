#!/usr/bin/env python3
"""
Agrega manual de uso y prepara espacios para imágenes al documento Word
"""

from docx import Document
from docx.shared import Pt, RGBColor, Inches, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

def agregar_manual_uso(doc):
    """Agrega sección completa de manual de uso"""
    
    doc.add_page_break()
    doc.add_heading("6. MANUAL DE USO DEL SIMULADOR", 1)
    
    doc.add_heading("6.1 Inicio Rápido", 2)
    p = doc.add_paragraph()
    p.add_run("PASO 1: ").bold = True
    p.add_run("Ejecutar el simulador").bold = False
    doc.add_paragraph("   • Abrir una terminal en el directorio del proyecto", style='List Bullet 2')
    doc.add_paragraph("   • Ejecutar el comando: python main.py", style='List Bullet 2')
    doc.add_paragraph("   • Se abrirá la ventana principal del simulador", style='List Bullet 2')
    doc.add_paragraph()
    
    p = doc.add_paragraph()
    p.add_run("PASO 2: ").bold = True
    p.add_run("Configurar parámetros iniciales").bold = False
    doc.add_paragraph("   • Establecer tamaño de memoria (ej: 1024 KB)", style='List Bullet 2')
    doc.add_paragraph("   • Seleccionar algoritmo de planificación", style='List Bullet 2')
    doc.add_paragraph("   • Si es Round Robin, configurar quantum (ej: 3)", style='List Bullet 2')
    doc.add_paragraph("   • Seleccionar estrategia de asignación de memoria", style='List Bullet 2')
    doc.add_paragraph()
    
    p = doc.add_paragraph()
    p.add_run("PASO 3: ").bold = True
    p.add_run("Crear procesos").bold = False
    doc.add_paragraph("   • Opción A: Agregar proceso manual (botón 'Agregar Proceso Manual')", style='List Bullet 2')
    doc.add_paragraph("   • Opción B: Generar test automático (botón 'Generar Test Automático')", style='List Bullet 2')
    doc.add_paragraph()
    
    p = doc.add_paragraph()
    p.add_run("PASO 4: ").bold = True
    p.add_run("Iniciar simulación").bold = False
    doc.add_paragraph("   • Presionar el botón 'INICIAR SIMULACIÓN'", style='List Bullet 2')
    doc.add_paragraph("   • Observar cómo los procesos se cargan a memoria", style='List Bullet 2')
    doc.add_paragraph("   • Ver cambios en tiempo real en todos los paneles", style='List Bullet 2')
    
    doc.add_heading("6.2 Descripción de Componentes de la Interfaz", 2)
    
    doc.add_heading("6.2.1 Panel de Configuración (Lado Izquierdo)", 3)
    doc.add_paragraph("El panel de configuración permite ajustar todos los parámetros del sistema:")
    
    config_items = [
        ("💾 Total RAM (KB)", "Campo numérico para establecer el tamaño total de memoria. Presionar 'Cambiar' para aplicar."),
        ("🔄 Algoritmo de Planificación", "ComboBox con 4 opciones: Round Robin, FCFS, SJF, Prioridad. El cambio se aplica inmediatamente."),
        ("⏱️ Quantum", "Solo visible para Round Robin. Establece el tiempo de ejecución antes del cambio de contexto."),
        ("🧩 Estrategia de Asignación", "ComboBox: First Fit, Best Fit, Worst Fit. Afecta cómo se asignan nuevos procesos."),
        ("➕ Agregar Proceso Manual", "Abre ventana modal para crear procesos personalizados con tamaño, tiempo y prioridad."),
        ("🎲 Generar Test Automático", "Crea 4 procesos de prueba predefinidos para demostración rápida."),
        ("▶ INICIAR SIMULACIÓN", "Botón principal que inicia/pausa la simulación. Cambia a '⏸ PAUSAR' durante la ejecución.")
    ]
    
    for item, desc in config_items:
        p = doc.add_paragraph()
        p.add_run(f"{item}: ").bold = True
        p.add_run(desc).bold = False
    
    doc.add_heading("6.2.2 Tabla de Procesos (Centro)", 3)
    doc.add_paragraph("Muestra el estado actual de todos los procesos en el sistema:")
    doc.add_paragraph("• PID: Identificador único del proceso", style='List Bullet')
    doc.add_paragraph("• Estado: Estado actual con iconos (🆕 NUEVO, ✅ LISTO, ⚡ EJECUCION, ⏸️ BLOQUEADO, ✔️ TERMINADO)", style='List Bullet')
    doc.add_paragraph("• Burst: Tiempo de ejecución restante en segundos", style='List Bullet')
    doc.add_paragraph("• Size: Tamaño del proceso en KB", style='List Bullet')
    doc.add_paragraph("• Prio: Prioridad del proceso (menor número = mayor prioridad)", style='List Bullet')
    doc.add_paragraph()
    doc.add_paragraph("La tabla se actualiza automáticamente cada 200ms durante la simulación.")
    
    doc.add_heading("6.2.3 Visualización de Memoria (Lado Derecho)", 3)
    doc.add_paragraph("Representación gráfica del mapa de memoria:")
    doc.add_paragraph("• Cada bloque se muestra como un rectángulo proporcional a su tamaño", style='List Bullet')
    doc.add_paragraph("• Bloques ocupados tienen colores únicos por proceso", style='List Bullet')
    doc.add_paragraph("• Bloques libres aparecen en gris", style='List Bullet')
    doc.add_paragraph("• El proceso en ejecución se resalta en dorado con borde brillante", style='List Bullet')
    doc.add_paragraph("• Al pasar el mouse sobre un bloque, se muestra información detallada abajo", style='List Bullet')
    doc.add_paragraph("• Las direcciones de memoria se muestran a la izquierda de cada bloque", style='List Bullet')
    
    doc.add_heading("6.2.4 Gráfico de Gantt (Abajo)", 3)
    doc.add_paragraph("Línea de tiempo que muestra el historial de uso de CPU:")
    doc.add_paragraph("• Cada segmento representa un período de ejecución", style='List Bullet')
    doc.add_paragraph("• Los colores corresponden a los procesos", style='List Bullet')
    doc.add_paragraph("• Muestra los últimos 60 ticks de ejecución", style='List Bullet')
    doc.add_paragraph("• El proceso actual se resalta en dorado", style='List Bullet')
    doc.add_paragraph("• Botón 'Limpiar Gantt' para reiniciar el historial", style='List Bullet')
    
    doc.add_heading("6.2.5 Log de Eventos (Abajo)", 3)
    doc.add_paragraph("Registro textual de todas las operaciones del sistema:")
    doc.add_paragraph("• Muestra eventos con timestamp (HH:MM:SS)", style='List Bullet')
    doc.add_paragraph("• Scroll automático al final", style='List Bullet')
    doc.add_paragraph("• Límite de 200 líneas (elimina las antiguas automáticamente)", style='List Bullet')
    doc.add_paragraph("• Fuente monospace para mejor legibilidad", style='List Bullet')
    
    doc.add_heading("6.3 Operaciones Básicas", 2)
    
    doc.add_heading("6.3.1 Agregar un Proceso Manualmente", 3)
    pasos_manual = [
        "Hacer clic en el botón '➕ Agregar Proceso Manual'",
        "Se abrirá una ventana modal con tres campos:",
        "  - Tamaño (KB): Ingresar el tamaño del proceso en kilobytes",
        "  - Tiempo de Ejecución (Seg): Duración total del proceso",
        "  - Prioridad: Nivel de prioridad (número entero, menor = mayor prioridad)",
        "Hacer clic en '✓ Confirmar y Agregar'",
        "El proceso aparecerá en la tabla en estado NUEVO",
        "Cuando haya memoria disponible, se cargará automáticamente"
    ]
    for paso in pasos_manual:
        doc.add_paragraph(paso, style='List Bullet')
    
    doc.add_heading("6.3.2 Cambiar Algoritmo de Planificación", 3)
    doc.add_paragraph("1. Seleccionar nuevo algoritmo en el ComboBox '🔄 Algoritmo de Planificación'")
    doc.add_paragraph("2. El cambio se aplica inmediatamente")
    doc.add_paragraph("3. Si se selecciona Round Robin, el campo Quantum se habilita")
    doc.add_paragraph("4. Si se selecciona otro algoritmo, el campo Quantum se deshabilita")
    doc.add_paragraph("5. Los procesos existentes se reordenan según el nuevo algoritmo")
    
    doc.add_heading("6.3.3 Cambiar Estrategia de Memoria", 3)
    doc.add_paragraph("1. Seleccionar nueva estrategia en el ComboBox '🧩 Estrategia de Asignación'")
    doc.add_paragraph("2. El cambio afecta solo a los nuevos procesos que se agreguen")
    doc.add_paragraph("3. Los procesos ya en memoria no se mueven")
    doc.add_paragraph("4. Las estrategias disponibles son:")
    doc.add_paragraph("   • First Fit: Asigna el primer bloque que cumpla", style='List Bullet 2')
    doc.add_paragraph("   • Best Fit: Asigna el bloque más pequeño que cumpla", style='List Bullet 2')
    doc.add_paragraph("   • Worst Fit: Asigna el bloque más grande disponible", style='List Bullet 2')
    
    doc.add_heading("6.3.4 Reiniciar el Sistema", 3)
    doc.add_paragraph("1. Cambiar el valor en '💾 Total RAM (KB)'")
    doc.add_paragraph("2. Presionar el botón '🔄 Cambiar'")
    doc.add_paragraph("3. El sistema se reinicia completamente:")
    doc.add_paragraph("   • Todos los procesos se eliminan", style='List Bullet 2')
    doc.add_paragraph("   • La memoria se reinicia con el nuevo tamaño", style='List Bullet 2')
    doc.add_paragraph("   • El Gantt se limpia", style='List Bullet 2')
    doc.add_paragraph("   • El log muestra el mensaje de reinicio", style='List Bullet 2')
    
    doc.add_heading("6.4 Interpretación de la Visualización", 2)
    
    doc.add_heading("6.4.1 Estados de Proceso", 3)
    estados_info = [
        ("🆕 NUEVO", "Proceso recién creado, esperando memoria disponible"),
        ("✅ LISTO", "Proceso en memoria, listo para ejecutarse, esperando CPU"),
        ("⚡ EJECUCION", "Proceso actualmente ejecutándose en la CPU"),
        ("⏸️ BLOQUEADO", "Proceso esperando I/O o algún recurso"),
        ("✔️ TERMINADO", "Proceso que ha finalizado, su memoria será liberada")
    ]
    for estado, desc in estados_info:
        p = doc.add_paragraph()
        p.add_run(f"{estado}: ").bold = True
        p.add_run(desc).bold = False
    
    doc.add_heading("6.4.2 Visualización de Memoria", 3)
    doc.add_paragraph("• Los bloques de memoria se muestran de arriba hacia abajo")
    doc.add_paragraph("• El tamaño de cada bloque es proporcional a su tamaño real")
    doc.add_paragraph("• El proceso en ejecución tiene un resaltado dorado brillante")
    doc.add_paragraph("• Al pasar el mouse, el label inferior muestra: estado, PID, tamaño y dirección")
    doc.add_paragraph("• Los bloques libres aparecen en gris y pueden ser reutilizados")
    
    doc.add_heading("6.4.3 Gráfico de Gantt", 3)
    doc.add_paragraph("• El gráfico muestra el historial de ejecución de izquierda a derecha")
    doc.add_paragraph("• Cada segmento representa un tick de ejecución")
    doc.add_paragraph("• Los segmentos del mismo color pertenecen al mismo proceso")
    doc.add_paragraph("• El segmento actual (dorado) muestra qué proceso está ejecutándose ahora")
    doc.add_paragraph("• El ancho de cada segmento es proporcional a su duración")
    
    doc.add_heading("6.5 Consejos y Mejores Prácticas", 2)
    consejos = [
        "Usa 'Generar Test Automático' para ver una demostración rápida del sistema",
        "Observa el log de eventos para entender el flujo completo del sistema",
        "Pasa el mouse sobre la memoria para ver detalles específicos de cada bloque",
        "Compara diferentes algoritmos de planificación con los mismos procesos",
        "Experimenta con diferentes estrategias de memoria para ver cómo afectan la asignación",
        "Usa procesos con diferentes tamaños y prioridades para ver el comportamiento",
        "El proceso en ejecución siempre se resalta en dorado en memoria y Gantt",
        "Puedes cambiar el algoritmo durante la ejecución para ver diferencias en tiempo real"
    ]
    for consejo in consejos:
        doc.add_paragraph(consejo, style='List Bullet')
    
    doc.add_heading("6.6 Solución de Problemas", 2)
    problemas = [
        ("Proceso no se carga a memoria", "Verificar que haya suficiente memoria disponible. El proceso permanecerá en estado NUEVO hasta que haya espacio."),
        ("No veo cambios en la interfaz", "Asegurarse de que la simulación esté iniciada (botón debe decir '⏸ PAUSAR')."),
        ("Quantum no funciona", "Verificar que el algoritmo seleccionado sea 'Round Robin'. El quantum solo aplica a este algoritmo."),
        ("Memoria no se actualiza", "La memoria se actualiza automáticamente. Si no hay cambios, verificar que los procesos estén en ejecución."),
        ("Gantt está vacío", "El Gantt solo muestra ejecuciones. Esperar a que los procesos comiencen a ejecutarse.")
    ]
    for problema, solucion in problemas:
        p = doc.add_paragraph()
        p.add_run(f"Problema: {problema}").bold = True
        doc.add_paragraph(f"   Solución: {solucion}", style='List Bullet 2')
    
    doc.add_page_break()

def agregar_seccion_imagenes(doc):
    """Agrega sección con información sobre las imágenes necesarias"""
    
    doc.add_heading("7. IMÁGENES DE LA INTERFAZ", 1)
    
    doc.add_paragraph("A continuación se listan las imágenes que deben capturarse de la interfaz para completar la documentación:")
    doc.add_paragraph()
    
    imagenes_necesarias = [
        {
            "nombre": "imagen_01_interfaz_completa.png",
            "descripcion": "Vista completa de la interfaz del simulador mostrando todos los paneles",
            "instrucciones": "Capturar la pantalla completa cuando la interfaz está abierta, mostrando panel de configuración, tabla de procesos, memoria, Gantt y log"
        },
        {
            "nombre": "imagen_02_panel_configuracion.png",
            "descripcion": "Panel de configuración detallado",
            "instrucciones": "Zoom al panel izquierdo mostrando todos los controles: memoria, algoritmo, quantum, estrategia y botones"
        },
        {
            "nombre": "imagen_03_tabla_procesos.png",
            "descripcion": "Tabla de procesos con varios procesos en diferentes estados",
            "instrucciones": "Capturar la tabla de procesos mostrando procesos en diferentes estados (NUEVO, LISTO, EJECUCION, BLOQUEADO)"
        },
        {
            "nombre": "imagen_04_memoria_ocupada.png",
            "descripcion": "Visualización de memoria con procesos asignados",
            "instrucciones": "Capturar la visualización de memoria cuando hay varios procesos cargados, mostrando bloques de diferentes colores"
        },
        {
            "nombre": "imagen_05_proceso_ejecutando.png",
            "descripcion": "Memoria mostrando proceso en ejecución (resaltado dorado)",
            "instrucciones": "Capturar cuando hay un proceso en ejecución, mostrando el resaltado dorado brillante del proceso activo"
        },
        {
            "nombre": "imagen_06_gantt_ejecutando.png",
            "descripcion": "Gráfico de Gantt con historial de ejecución",
            "instrucciones": "Capturar el gráfico de Gantt cuando hay varios ticks de ejecución, mostrando diferentes colores por proceso"
        },
        {
            "nombre": "imagen_07_log_eventos.png",
            "descripcion": "Log de eventos mostrando mensajes del sistema",
            "instrucciones": "Capturar el log con varios mensajes visibles, mostrando timestamps y diferentes tipos de eventos"
        },
        {
            "nombre": "imagen_08_ventana_agregar_proceso.png",
            "descripcion": "Ventana modal para agregar proceso manualmente",
            "instrucciones": "Abrir la ventana de agregar proceso y capturar la ventana modal completa"
        },
        {
            "nombre": "imagen_09_hover_memoria.png",
            "descripcion": "Hover informativo sobre bloque de memoria",
            "instrucciones": "Pasar el mouse sobre un bloque de memoria y capturar cuando se muestra la información detallada en el label inferior"
        },
        {
            "nombre": "imagen_10_comparacion_algoritmos.png",
            "descripcion": "Comparación visual de diferentes algoritmos",
            "instrucciones": "Capturar la interfaz con Round Robin y luego con SJF, mostrando diferencias en el orden de ejecución"
        }
    ]
    
    doc.add_heading("7.1 Lista de Imágenes Requeridas", 2)
    
    tabla = doc.add_table(rows=1, cols=3)
    tabla.style = 'Light Grid Accent 1'
    
    # Encabezado
    hdr = tabla.rows[0].cells
    hdr[0].text = 'Nombre del Archivo'
    hdr[1].text = 'Descripción'
    hdr[2].text = 'Instrucciones de Captura'
    for cell in hdr:
        cell.paragraphs[0].runs[0].font.bold = True
        shading_elm = OxmlElement('w:shd')
        shading_elm.set(qn('w:fill'), '4472C4')
        cell._element.get_or_add_tcPr().append(shading_elm)
        cell.paragraphs[0].runs[0].font.color.rgb = RGBColor(255, 255, 255)
    
    # Datos
    for img in imagenes_necesarias:
        row = tabla.add_row().cells
        row[0].text = img['nombre']
        row[0].paragraphs[0].runs[0].font.bold = True
        row[1].text = img['descripcion']
        row[2].text = img['instrucciones']
    
    doc.add_paragraph()
    
    doc.add_heading("7.2 Instrucciones para Capturar Imágenes", 2)
    doc.add_paragraph("Para capturar las imágenes de la interfaz:")
    doc.add_paragraph("1. Ejecutar el simulador: python main.py", style='List Number')
    doc.add_paragraph("2. Configurar parámetros iniciales (memoria, algoritmo, etc.)", style='List Number')
    doc.add_paragraph("3. Crear algunos procesos (manual o test automático)", style='List Number')
    doc.add_paragraph("4. Iniciar la simulación", style='List Number')
    doc.add_paragraph("5. Usar herramienta de captura de pantalla (Snipping Tool, Print Screen, etc.)", style='List Number')
    doc.add_paragraph("6. Guardar las imágenes con los nombres especificados en formato PNG", style='List Number')
    doc.add_paragraph("7. Insertar las imágenes en el documento Word en las secciones correspondientes", style='List Number')
    
    doc.add_paragraph()
    doc.add_paragraph("NOTA: Las imágenes deben tener buena resolución y mostrar claramente los elementos de la interfaz. Se recomienda usar formato PNG para mantener calidad.")

def actualizar_documento():
    """Actualiza el documento Word existente agregando manual de uso"""
    
    doc = Document('PRESENTACION_PROYECTO.docx')
    
    # Agregar manual de uso
    agregar_manual_uso(doc)
    
    # Agregar sección de imágenes
    agregar_seccion_imagenes(doc)
    
    # Guardar
    doc.save('PRESENTACION_PROYECTO.docx')
    print("✅ Manual de uso agregado al documento")
    print(f"   - Total párrafos: {len(doc.paragraphs)}")
    print(f"   - Total tablas: {len(doc.tables)}")
    print("   - Sección 6: Manual de Uso completo")
    print("   - Sección 7: Lista de imágenes requeridas")

if __name__ == '__main__':
    actualizar_documento()
