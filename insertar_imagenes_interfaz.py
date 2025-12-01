#!/usr/bin/env python3
"""
Inserta las imágenes en la sección 3 de interfaz y elimina la sección 7
"""

from docx import Document
from docx.shared import Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
import os
import glob

def insertar_imagenes_en_interfaz(doc):
    """Inserta las imágenes en la sección 3.2"""
    imagenes_path = "Imagenes/"
    imagenes = sorted(glob.glob(os.path.join(imagenes_path, "*.png")))
    
    if not imagenes:
        print("⚠️ No se encontraron imágenes")
        return False
    
    print(f"📸 Insertando {len(imagenes)} imágenes...")
    
    # Buscar la sección 3.2.5 (Log de Eventos) para insertar después
    indice_insertar = -1
    for i, para in enumerate(doc.paragraphs):
        texto = para.text.strip()
        if "3.2.5 Log de Eventos" in texto or "6.2.5 Log de Eventos" in texto:
            # Buscar el final de esta subsección (antes de 3.3 o 3.4)
            j = i + 1
            while j < len(doc.paragraphs):
                texto_sig = doc.paragraphs[j].text.strip()
                if texto_sig.startswith("3.3") or texto_sig.startswith("3.4") or \
                   texto_sig.startswith("6.3") or texto_sig.startswith("6.4"):
                    indice_insertar = j
                    break
                j += 1
            break
    
    if indice_insertar == -1:
        # Buscar "3.3" o "3.4" directamente
        for i, para in enumerate(doc.paragraphs):
            texto = para.text.strip()
            if "3.3 Ventana" in texto or "3.4 Sistema" in texto or \
               "6.3 Operaciones" in texto or "6.4 Interpretación" in texto:
                indice_insertar = i
                break
    
    if indice_insertar == -1:
        print("⚠️ No se encontró dónde insertar las imágenes")
        return False
    
    # Insertar título de subsección
    para_antes = doc.paragraphs[indice_insertar]
    para_antes.insert_paragraph_before("")
    
    heading = doc.paragraphs[indice_insertar]
    heading.insert_paragraph_before("3.3 Capturas de Pantalla de la Interfaz")
    doc.paragraphs[indice_insertar].style = 'Heading 3'
    
    # Descripción
    doc.paragraphs[indice_insertar].insert_paragraph_before(
        "A continuación se presentan capturas de pantalla de los diferentes componentes de la interfaz del simulador:")
    
    # Descripciones de las imágenes
    descripciones = [
        "Vista completa de la interfaz mostrando todos los paneles: configuración, tabla de procesos, visualización de memoria, gráfico de Gantt y log de eventos.",
        "Panel de configuración detallado con todos los controles: configuración de memoria, selección de algoritmo, quantum, estrategia de memoria y botones de acción.",
        "Tabla de procesos mostrando diferentes estados: procesos en estado NUEVO, LISTO, EJECUCION y BLOQUEADO con sus respectivos iconos.",
        "Visualización gráfica del mapa de memoria con procesos asignados, mostrando bloques de diferentes colores y el proceso en ejecución resaltado en dorado.",
        "Gráfico de Gantt mostrando el historial de ejecución de procesos con diferentes colores por proceso y el segmento actual resaltado.",
        "Log de eventos con mensajes del sistema mostrando timestamps y diferentes tipos de eventos registrados durante la simulación."
    ]
    
    # Insertar cada imagen
    for idx, (img_path, desc) in enumerate(zip(imagenes, descripciones), 1):
        # Título de la figura
        doc.paragraphs[indice_insertar].insert_paragraph_before("")
        p_titulo = doc.paragraphs[indice_insertar]
        p_titulo.insert_paragraph_before(f"Figura {idx}: {desc}")
        doc.paragraphs[indice_insertar].runs[0].bold = True
        doc.paragraphs[indice_insertar].runs[0].font.size = Inches(0.12)
        
        # Insertar imagen centrada
        doc.paragraphs[indice_insertar].insert_paragraph_before("")
        p_img = doc.paragraphs[indice_insertar]
        p_img.insert_paragraph_before("")
        p_img = doc.paragraphs[indice_insertar]
        p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
        
        try:
            run = p_img.add_run()
            # Ajustar tamaño según la imagen
            if idx == 1:  # Interfaz completa - más grande
                run.add_picture(img_path, width=Inches(7))
            else:
                run.add_picture(img_path, width=Inches(6))
            print(f"✅ Imagen {idx} agregada: {os.path.basename(img_path)}")
        except Exception as e:
            print(f"⚠️ Error al agregar imagen {img_path}: {e}")
            p_img.add_run(f"[Imagen no disponible: {os.path.basename(img_path)}]")
        
        doc.paragraphs[indice_insertar].insert_paragraph_before("")
    
    return True

def eliminar_seccion_7(doc):
    """Elimina completamente la sección 7"""
    indices_eliminar = []
    en_seccion_7 = False
    
    for i, para in enumerate(doc.paragraphs):
        texto = para.text.strip().upper()
        if "7. IMÁGENES" in texto or "7. IMAGENES" in texto:
            en_seccion_7 = True
        
        if en_seccion_7:
            # Continuar eliminando hasta encontrar sección 8 o fin del documento
            if texto.startswith("8.") or (i + 1 < len(doc.paragraphs) and 
                (doc.paragraphs[i+1].text.strip().startswith("8.") or 
                 doc.paragraphs[i+1].text.strip().startswith("CONCLUSIÓN"))):
                break
            indices_eliminar.append(i)
    
    # Eliminar en orden inverso
    eliminados = 0
    for i in reversed(indices_eliminar):
        try:
            p = doc.paragraphs[i]._element
            p.getparent().remove(p)
            eliminados += 1
        except:
            pass
    
    print(f"✅ Sección 7 eliminada ({eliminados} párrafos)")
    return eliminados

def actualizar_documento():
    """Actualiza el documento"""
    doc = Document('PRESENTACION_PROYECTO.docx')
    
    print("📝 Actualizando documento Word...")
    print(f"   Párrafos iniciales: {len(doc.paragraphs)}")
    
    # Eliminar sección 7 primero
    eliminar_seccion_7(doc)
    
    # Insertar imágenes en sección 3
    if insertar_imagenes_en_interfaz(doc):
        print("✅ Imágenes insertadas en sección 3.3")
    
    # Guardar
    doc.save('PRESENTACION_PROYECTO.docx')
    print(f"✅ Documento guardado")
    print(f"   - Párrafos finales: {len(doc.paragraphs)}")
    
    # Verificar imágenes
    imagenes_en_doc = 0
    for para in doc.paragraphs:
        for run in para.runs:
            if run._element.xpath('.//a:blip'):
                imagenes_en_doc += 1
    
    print(f"   - Imágenes en documento: {imagenes_en_doc}")

if __name__ == '__main__':
    actualizar_documento()
