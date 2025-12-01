#!/usr/bin/env python3
"""
Elimina definitivamente la sección 7 del documento
"""

from docx import Document

def eliminar_seccion_7_completa(doc):
    """Elimina la sección 7 completamente"""
    para_eliminar = []
    en_seccion_7 = False
    
    for i, para in enumerate(doc.paragraphs):
        texto = para.text.strip().upper()
        
        # Detectar inicio de sección 7
        if "7. IM" in texto and ("ÁGENES" in texto or "AGENES" in texto):
            en_seccion_7 = True
            print(f"📍 Inicio sección 7 encontrado en línea {i}: {para.text[:60]}")
        
        # Si estamos en sección 7, marcar para eliminar
        if en_seccion_7:
            # Continuar hasta encontrar sección 8 o fin
            if texto.startswith("8.") or (i + 1 < len(doc.paragraphs) and 
                doc.paragraphs[i+1].text.strip().startswith("8.")):
                print(f"📍 Fin de sección 7 en línea {i}")
                break
            para_eliminar.append(i)
    
    # Eliminar en orden inverso
    eliminados = 0
    for i in reversed(para_eliminar):
        try:
            p = doc.paragraphs[i]._element
            p.getparent().remove(p)
            eliminados += 1
        except Exception as e:
            print(f"⚠️ Error eliminando línea {i}: {e}")
    
    return eliminados

# Cargar y actualizar
doc = Document('PRESENTACION_PROYECTO.docx')
print(f"📝 Documento cargado: {len(doc.paragraphs)} párrafos")

eliminados = eliminar_seccion_7_completa(doc)

if eliminados > 0:
    doc.save('PRESENTACION_PROYECTO.docx')
    print(f"✅ Sección 7 eliminada: {eliminados} párrafos")
    print(f"   Párrafos finales: {len(doc.paragraphs)}")
else:
    print("ℹ️ Sección 7 no encontrada o ya eliminada")
