#!/usr/bin/env python3
"""
Script para crear un documento Word profesional con tablas, formato y estructura mejorada
"""

from docx import Document
from docx.shared import Pt, RGBColor, Inches, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import re

def agregar_titulo_principal(doc, texto):
    """Agrega un título principal con formato"""
    p = doc.add_heading(texto, level=0)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in p.runs:
        run.font.size = Pt(20)
        run.font.bold = True
        run.font.color.rgb = RGBColor(0, 51, 102)
    return p

def agregar_seccion(doc, titulo, nivel=1):
    """Agrega una sección con formato"""
    heading = doc.add_heading(titulo, level=nivel)
    for run in heading.runs:
        run.font.color.rgb = RGBColor(0, 102, 204)
    doc.add_paragraph()  # Espacio después del título
    return heading

def crear_tabla_requerimientos(doc):
    """Crea una tabla profesional de requerimientos"""
    doc.add_paragraph("A continuación se presenta un resumen de los requerimientos implementados:")
    doc.add_paragraph()
    
    # Tabla de Gestión de Procesos
    tabla = doc.add_table(rows=1, cols=2)
    tabla.style = 'Light Grid Accent 1'
    tabla.alignment = WD_TABLE_ALIGNMENT.CENTER
    
    # Encabezado
    hdr_cells = tabla.rows[0].cells
    hdr_cells[0].text = 'Componente'
    hdr_cells[1].text = 'Funcionalidad Implementada'
    for cell in hdr_cells:
        cell.paragraphs[0].runs[0].font.bold = True
        cell.paragraphs[0].runs[0].font.color.rgb = RGBColor(255, 255, 255)
        shading_elm = OxmlElement('w:shd')
        shading_elm.set(qn('w:fill'), '4472C4')
        cell._element.get_or_add_tcPr().append(shading_elm)
    
    # Datos
    datos = [
        ('Gestión de Procesos', 'Creación de procesos, 5 estados clásicos, Colas de procesos, PCB, Transiciones automáticas'),
        ('Gestión de Memoria', 'Asignación dinámica, 3 algoritmos (First/Best/Worst Fit), Liberación automática, Compactación, Visualización en tiempo real'),
        ('Planificación', '4 algoritmos: Round Robin, FCFS, SJF, Prioridad. Cambio dinámico, Quantum configurable, Gráfico de Gantt'),
        ('CPU', 'Ejecución por ticks, Manejo de quantum, Bloqueos por I/O, Retorno automático'),
        ('Interfaz Gráfica', 'Panel de configuración, Tabla de procesos, Visualización de memoria, Gráfico de Gantt, Log de eventos, Tema oscuro futurista')
    ]
    
    for componente, funcionalidad in datos:
        row_cells = tabla.add_row().cells
        row_cells[0].text = componente
        row_cells[0].paragraphs[0].runs[0].font.bold = True
        row_cells[1].text = funcionalidad
    
    doc.add_paragraph()

def crear_tabla_arquitectura(doc):
    """Crea tabla de arquitectura del sistema"""
    doc.add_paragraph("Estructura modular del sistema:")
    doc.add_paragraph()
    
    tabla = doc.add_table(rows=1, cols=3)
    tabla.style = 'Medium Grid 1 Accent 1'
    
    # Encabezado
    hdr = tabla.rows[0].cells
    hdr[0].text = 'Módulo'
    hdr[1].text = 'Archivo'
    hdr[2].text = 'Responsabilidad'
    for cell in hdr:
        cell.paragraphs[0].runs[0].font.bold = True
        shading_elm = OxmlElement('w:shd')
        shading_elm.set(qn('w:fill'), 'D9E1F2')
        cell._element.get_or_add_tcPr().append(shading_elm)
    
    modulos = [
        ('Interfaz', 'interfaz.py', 'Interfaz gráfica y visualización en tiempo real'),
        ('Coordinador', 'coordinador.py', 'Integración y coordinación de todos los módulos'),
        ('CPU', 'modulo_cpu.py', 'Ejecución de procesos por ticks'),
        ('Procesos', 'modulo_procesos.py', 'Gestión de procesos y estados'),
        ('Memoria', 'modulo_memoria.py', 'Asignación y gestión de memoria'),
        ('Planificador', 'modulo_planificador.py', 'Selección de procesos a ejecutar'),
        ('Despachador', 'modulo_despachador.py', 'Asignación de CPU a procesos'),
        ('Constantes', 'constantes.py', 'Configuraciones y constantes del sistema')
    ]
    
    for modulo, archivo, responsabilidad in modulos:
        row = tabla.add_row().cells
        row[0].text = modulo
        row[0].paragraphs[0].runs[0].font.bold = True
        row[1].text = archivo
        row[2].text = responsabilidad
    
    doc.add_paragraph()

def crear_tabla_componentes_interfaz(doc):
    """Crea tabla detallada de componentes de la interfaz"""
    doc.add_paragraph("Componentes principales de la interfaz gráfica:")
    doc.add_paragraph()
    
    tabla = doc.add_table(rows=1, cols=2)
    tabla.style = 'Light List Accent 1'
    
    hdr = tabla.rows[0].cells
    hdr[0].text = 'Componente'
    hdr[1].text = 'Descripción y Funcionalidades'
    for cell in hdr:
        cell.paragraphs[0].runs[0].font.bold = True
        shading_elm = OxmlElement('w:shd')
        shading_elm.set(qn('w:fill'), 'E7E6E6')
        cell._element.get_or_add_tcPr().append(shading_elm)
    
    componentes = [
        ('Panel de Configuración', 
         '• Configuración de memoria total (KB)\n• Selección de algoritmo de planificación\n• Ajuste de quantum (Round Robin)\n• Estrategia de asignación de memoria\n• Botones de acción (Agregar proceso, Test automático, Iniciar/Pausar)'),
        ('Tabla de Procesos',
         '• Muestra todos los procesos con estados actualizados\n• Columnas: PID, Estado (con iconos), Burst restante, Tamaño, Prioridad\n• Actualización en tiempo real\n• Resaltado del proceso en ejecución'),
        ('Visualización de Memoria',
         '• Representación gráfica del mapa de memoria\n• Bloques con colores únicos por proceso\n• Resaltado dorado del proceso en ejecución\n• Hover informativo con detalles (PID, tamaño, dirección)\n• Escalado automático'),
        ('Gráfico de Gantt',
         '• Línea de tiempo de uso de CPU\n• Muestra últimos 60 ticks\n• Colores corresponden a procesos\n• Resaltado del proceso actual\n• Botón para limpiar historial'),
        ('Log de Eventos',
         '• Registro textual de todas las operaciones\n• Timestamps (HH:MM:SS)\n• Scroll automático\n• Límite de 200 líneas\n• Fuente monospace para legibilidad')
    ]
    
    for componente, descripcion in componentes:
        row = tabla.add_row().cells
        row[0].text = componente
        row[0].paragraphs[0].runs[0].font.bold = True
        row[1].text = descripcion
    
    doc.add_paragraph()

def crear_tabla_algoritmos(doc):
    """Crea tabla de algoritmos implementados"""
    doc.add_paragraph("Algoritmos de planificación implementados:")
    doc.add_paragraph()
    
    tabla = doc.add_table(rows=1, cols=3)
    tabla.style = 'Colorful Grid Accent 2'
    
    hdr = tabla.rows[0].cells
    hdr[0].text = 'Algoritmo'
    hdr[1].text = 'Descripción'
    hdr[2].text = 'Características'
    for cell in hdr:
        cell.paragraphs[0].runs[0].font.bold = True
    
    algoritmos = [
        ('Round Robin', 'FIFO con quantum', 'Quantum configurable, Cambio de contexto por tiempo'),
        ('FCFS', 'First Come First Served', 'FIFO simple, Sin preemption'),
        ('SJF', 'Shortest Job First', 'Ordena por tiempo de ejecución, Minimiza tiempo de espera'),
        ('Prioridad', 'Por nivel de prioridad', 'Menor número = mayor prioridad, Puede ser preemptivo')
    ]
    
    for algoritmo, desc, caracteristicas in algoritmos:
        row = tabla.add_row().cells
        row[0].text = algoritmo
        row[0].paragraphs[0].runs[0].font.bold = True
        row[1].text = desc
        row[2].text = caracteristicas
    
    doc.add_paragraph()

def crear_tabla_estados_proceso(doc):
    """Crea tabla de estados de proceso"""
    doc.add_paragraph("Estados de proceso implementados:")
    doc.add_paragraph()
    
    tabla = doc.add_table(rows=1, cols=2)
    tabla.style = 'Medium Shading 1 Accent 1'
    
    hdr = tabla.rows[0].cells
    hdr[0].text = 'Estado'
    hdr[1].text = 'Descripción'
    for cell in hdr:
        cell.paragraphs[0].runs[0].font.bold = True
    
    estados = [
        ('NUEVO', 'Proceso recién creado, esperando ser cargado a memoria'),
        ('LISTO', 'Proceso en memoria, listo para ejecutarse, esperando CPU'),
        ('EJECUCION', 'Proceso actualmente ejecutándose en la CPU'),
        ('BLOQUEADO', 'Proceso esperando I/O o algún recurso'),
        ('TERMINADO', 'Proceso que ha finalizado su ejecución')
    ]
    
    for estado, descripcion in estados:
        row = tabla.add_row().cells
        row[0].text = estado
        row[0].paragraphs[0].runs[0].font.bold = True
        row[1].text = descripcion
    
    doc.add_paragraph()

def procesar_texto_con_formato(parrafo, texto):
    """Procesa texto Markdown y aplica formato"""
    partes = re.split(r'(\*\*[^*]+\*\*|`[^`]+`)', texto)
    
    for parte in partes:
        if parte.startswith('**') and parte.endswith('**'):
            run = parrafo.add_run(parte[2:-2])
            run.bold = True
        elif parte.startswith('`') and parte.endswith('`'):
            run = parrafo.add_run(parte[1:-1])
            run.font.name = 'Consolas'
            run.font.size = Pt(10)
        else:
            if parte.strip():
                parrafo.add_run(parte)

def crear_documento_profesional():
    """Crea un documento Word profesional"""
    
    # Leer contenido del Markdown
    with open('PRESENTACION_PROYECTO.md', 'r', encoding='utf-8') as f:
        contenido = f.read()
    
    # Crear documento
    doc = Document()
    
    # Configurar márgenes
    sections = doc.sections
    for section in sections:
        section.top_margin = Cm(2.5)
        section.bottom_margin = Cm(2.5)
        section.left_margin = Cm(2.5)
        section.right_margin = Cm(2.5)
    
    # PORTADA
    doc.add_paragraph()
    titulo = agregar_titulo_principal(doc, "SIMULADOR DE SISTEMA OPERATIVO")
    doc.add_paragraph()
    subtitulo = doc.add_paragraph("Gestión de Memoria y Procesos")
    subtitulo.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in subtitulo.runs:
        run.font.size = Pt(16)
        run.font.color.rgb = RGBColor(102, 102, 102)
    doc.add_page_break()
    
    # 1. ARQUITECTURA
    agregar_seccion(doc, "1. EXPLICACIÓN DE LA ARQUITECTURA ASUMIDA", 1)
    
    # 1.1 Tecnología Base
    doc.add_heading("1.1 Tecnología Base", level=2)
    tabla_tech = doc.add_table(rows=5, cols=2)
    tabla_tech.style = 'Light Grid Accent 1'
    tech_data = [
        ('Lenguaje de programación', 'Python 3'),
        ('Framework gráfico', 'Tkinter'),
        ('Paradigma', 'Programación orientada a objetos (POO)'),
        ('Arquitectura', 'Modular y desacoplada'),
        ('Patrón de diseño', 'Coordinador Central (Mediator Pattern)')
    ]
    for i, (campo, valor) in enumerate(tech_data):
        tabla_tech.rows[i].cells[0].text = campo
        tabla_tech.rows[i].cells[0].paragraphs[0].runs[0].font.bold = True
        tabla_tech.rows[i].cells[1].text = valor
    
    doc.add_paragraph()
    
    # 1.2 Arquitectura General
    doc.add_heading("1.2 Arquitectura General del Sistema", level=2)
    doc.add_paragraph("El sistema está estructurado en 8 módulos independientes que se comunican a través de un coordinador central:")
    doc.add_paragraph()
    crear_tabla_arquitectura(doc)
    
    # 2. REQUERIMIENTOS
    agregar_seccion(doc, "2. REQUERIMIENTOS ATENDIDOS", 1)
    crear_tabla_requerimientos(doc)
    
    # Tabla de algoritmos
    crear_tabla_algoritmos(doc)
    
    # Tabla de estados
    crear_tabla_estados_proceso(doc)
    
    # 3. INTERFAZ GRÁFICA
    agregar_seccion(doc, "3. DESARROLLO DE LA INTERFAZ GRÁFICA", 1)
    
    doc.add_heading("3.1 Arquitectura de la Interfaz", level=2)
    p = doc.add_paragraph()
    p.add_run("La interfaz gráfica está implementada utilizando ").bold = False
    p.add_run("Tkinter").bold = True
    p.add_run(" con un enfoque orientado a objetos. La clase principal ").bold = False
    p.add_run("SimuladorApp").bold = True
    p.add_run(" gestiona todos los componentes visuales.")
    
    doc.add_paragraph()
    info_interfaz = doc.add_table(rows=4, cols=2)
    info_interfaz.style = 'Light List Accent 2'
    interfaz_data = [
        ('Ventana principal', '1400x950 píxeles con tema oscuro futurista'),
        ('Layout', 'Componentes organizados en paneles independientes'),
        ('Actualización', 'Bucle de simulación que actualiza la UI cada 200ms'),
        ('Comunicación', 'Bidireccional: envía comandos y recibe actualizaciones')
    ]
    for i, (campo, valor) in enumerate(interfaz_data):
        info_interfaz.rows[i].cells[0].text = campo
        info_interfaz.rows[i].cells[0].paragraphs[0].runs[0].font.bold = True
        info_interfaz.rows[i].cells[1].text = valor
    
    doc.add_paragraph()
    
    doc.add_heading("3.2 Componentes Principales", level=2)
    crear_tabla_componentes_interfaz(doc)
    
    # 4. DEMOSTRACIÓN
    agregar_seccion(doc, "4. DEMOSTRACIÓN DEL MÓDULO DESARROLLADO", 1)
    
    doc.add_heading("4.1 Funcionamiento del Sistema", level=2)
    doc.add_paragraph("El simulador ejecuta ciclos continuos mientras la simulación está activa:")
    
    pasos = [
        "Carga de procesos nuevos a memoria",
        "Asignación de memoria según estrategia configurada",
        "Manejo de retorno de procesos bloqueados (I/O)",
        "Ejecución de un tick en CPU",
        "Planificación del próximo proceso",
        "Despacho y asignación de CPU",
        "Actualización de la interfaz en tiempo real"
    ]
    
    for i, paso in enumerate(pasos, 1):
        p = doc.add_paragraph(f"{i}. {paso}", style='List Number')
    
    # 5. RESUMEN
    agregar_seccion(doc, "5. RESUMEN EJECUTIVO", 1)
    
    doc.add_paragraph("Características destacadas del proyecto:")
    lista_caracteristicas = [
        "Arquitectura modular bien estructurada y mantenible",
        "4 algoritmos de planificación implementados y funcionales",
        "3 estrategias de memoria con visualización en tiempo real",
        "Interfaz gráfica completa con múltiples visualizaciones",
        "Simulación en tiempo real con control total del usuario",
        "Documentación completa de todos los módulos"
    ]
    
    for caracteristica in lista_caracteristicas:
        p = doc.add_paragraph(caracteristica, style='List Bullet')
        p.runs[0].font.size = Pt(11)
    
    # Guardar
    doc.save('PRESENTACION_PROYECTO.docx')
    print("✅ Documento Word profesional creado: PRESENTACION_PROYECTO.docx")
    print("   - Incluye tablas formateadas")
    print("   - Estructura profesional")
    print("   - Formato mejorado")

if __name__ == '__main__':
    crear_documento_profesional()
